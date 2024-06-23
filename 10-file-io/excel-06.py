import os
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, PieChart

dirname, filename = os.path.split(os.path.abspath(__file__))

# Create a new workbook and select the active worksheet
wb = Workbook()
ws = wb.active

# Add some data to the worksheet
data = [
    ["Item", "Amount"],
    ["Item 1", 10],
    ["Item 2", 20],
    ["Item 3", 30],
    ["Item 4", 40],
    ["Item 5", 50]
]

for row in data:
    ws.append(row)

# Create a bar chart ================================================================
chart = BarChart()

# Select the data for the chart
data = Reference(ws, min_col=2, min_row=1, max_col=2, max_row=6)
categories = Reference(ws, min_col=1, min_row=2, max_row=6)

# Add data and categories to the chart
chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)

# Set chart title and labels
chart.title = "Sample Bar Chart"
chart.x_axis.title = "Items"
chart.y_axis.title = "Amount"

# Add the chart to the worksheet
ws.add_chart(chart, "E5")
# Create a bar chart ================================================================
# Create a pie chart ================================================================
chart2 = PieChart()

# Select the data for the chart2
data = Reference(ws, min_col=2, min_row=1, max_col=2, max_row=6)
categories = Reference(ws, min_col=1, min_row=2, max_row=6)

# Add data and categories to the chart2
chart2.add_data(data, titles_from_data=True)
chart2.set_categories(categories)

# Set chart2 title and labels
chart2.title = "Sample Pie Chart"


# Add the chart2 to the worksheet
ws.add_chart(chart2, "E20")

# Create a pie chart ================================================================




# Save the workbook
wb.save(f"{dirname}/chart_example.xlsx")
